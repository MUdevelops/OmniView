#!/usr/bin/env python3
"""Generate every test fixture programmatically (no downloads).
Needs: Pillow, reportlab, pypdf, python-docx, openpyxl, python-pptx, ImageMagick `convert`, LibreOffice `soffice`.
Not producible here: HEIC (no HEVC encoder) -> tests/e2e/heic.spec.ts skips when tests/fixtures/sample.heic is absent. Drop a real one in to enable it.
RAW fixtures are SYNTHETIC (TIFF container + embedded JPEGs), not camera output."""
import io, json, os, shutil, struct, subprocess, sys, tempfile, zipfile
from pathlib import Path
from PIL import Image, ImageDraw

OUT = Path(__file__).resolve().parent.parent / 'tests' / 'fixtures'
OUT.mkdir(parents=True, exist_ok=True)
EXPECT = {}  # fixture -> expected detected kind
def put(name, data, kind):
    (OUT / name).write_bytes(data if isinstance(data, bytes) else data.encode()); EXPECT[name] = kind

def art(w, h, a=(30, 90, 220), b=(250, 190, 30)):
    im = Image.new('RGB', (w, h)); px = im.load()
    for y in range(h):
        for x in range(w):
            t = x / max(1, w - 1); px[x, y] = tuple(int(a[i] * (1 - t) + b[i] * t) for i in range(3))
    d = ImageDraw.Draw(im); d.ellipse([w * .25, h * .2, w * .75, h * .8], outline='white', width=max(2, w // 40)); d.rectangle([w * .4, h * .4, w * .6, h * .6], fill=(200, 30, 60))
    return im
def encode(im, fmt, **kw):
    b = io.BytesIO(); im.save(b, fmt, **kw); return b.getvalue()

# ---- native images ---------------------------------------------------------------------------
exif = Image.Exif(); exif[0x010F] = 'OmniView Test'; exif[0x0110] = 'Fixture Cam'; exif[0x0132] = '2024:01:02 03:04:05'
put('sample.jpg', encode(art(320, 240), 'JPEG', quality=90, exif=exif), 'jpeg')
t = Image.new('RGBA', (200, 200), (0, 0, 0, 0)); ImageDraw.Draw(t).ellipse([20, 20, 180, 180], fill=(220, 40, 70, 255)); ImageDraw.Draw(t).ellipse([70, 70, 130, 130], fill=(255, 255, 255, 128))
put('transparent.png', encode(t, 'PNG'), 'png')
frames = [art(120, 90, (200, 30, 30), (30, 30, 200)), art(120, 90, (30, 200, 30), (200, 30, 200)), art(120, 90, (30, 30, 200), (250, 250, 30))]
put('animated.gif', encode(frames[0], 'GIF', save_all=True, append_images=frames[1:], duration=200, loop=0), 'gif')
put('sample.webp', encode(art(256, 192), 'WEBP', quality=85), 'webp')
put('sample.avif', encode(art(256, 192), 'AVIF', quality=60), 'avif')
put('sample.bmp', encode(art(200, 150), 'BMP'), 'bmp')
SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="240" height="160" viewBox="0 0 240 160"><rect width="240" height="160" fill="#1f5eff"/><circle cx="120" cy="80" r="50" fill="#fff"/><text x="120" y="88" font-size="20" text-anchor="middle" fill="#1f5eff">SVG</text></svg>'
put('sample.svg', SVG, 'svg')
put('evil.svg', SVG.replace('<rect', '<script>window.__pwned=1</script><rect onload="window.__pwned=1"') , 'svg')

# ---- decoded images --------------------------------------------------------------------------
pages = [art(160, 120, (200, 30, 30), (30, 30, 200)), art(200, 150, (30, 200, 30), (200, 30, 200)), art(120, 160, (30, 30, 200), (250, 250, 30))]
put('multipage.tiff', encode(pages[0], 'TIFF', save_all=True, append_images=pages[1:]), 'tiff')
with tempfile.TemporaryDirectory() as tmp:
    psd = os.path.join(tmp, 'l.psd')
    subprocess.run(['convert', '-size', '200x150', 'xc:#87ceeb', '-set', 'label', 'Background',
                    '(', '-size', '120x80', 'xc:#e53935', '-set', 'label', 'Red box', '-repage', '+20+20', ')',
                    '(', '-size', '80x80', 'xc:#43a047', '-set', 'label', 'Green square', '-repage', '+90+50', ')', '-depth', '8', psd], check=True)
    put('layered.psd', Path(psd).read_bytes(), 'psd')

# Synthetic RAW: little-endian TIFF with Make/Model/Orientation(6 = rotate 90) and embedded JPEG previews (160x120 thumb + 1200x800 preview).
def fake_raw(with_preview):
    thumb = encode(art(160, 120), 'JPEG', quality=70); big = encode(art(1200, 800), 'JPEG', quality=85)
    n = 5 if with_preview else 3
    ifd_end = 8 + 2 + n * 12 + 4; make_off = ifd_end; model_off = make_off + 9; data_off = (model_off + 10 + 3) & ~3
    thumb_off = data_off; big_off = thumb_off + len(thumb)
    ents = [struct.pack('<HHII', 271, 2, 9, make_off), struct.pack('<HHII', 272, 2, 10, model_off), struct.pack('<HHIHH', 274, 3, 1, 6, 0)]
    if with_preview: ents += [struct.pack('<HHII', 513, 4, 1, thumb_off), struct.pack('<HHII', 514, 4, 1, len(thumb))]
    body = b'II*\x00' + struct.pack('<I', 8) + struct.pack('<H', n) + b''.join(ents) + struct.pack('<I', 0) + b'OmniView\0' + b'Synthetic\0'
    body += b'\0' * (data_off - len(body))
    if with_preview: body += thumb + big
    return body
put('photo.dng', fake_raw(True), 'raw'); put('photo-nopreview.dng', fake_raw(False), 'raw')

# ---- PDF / AI / EPS --------------------------------------------------------------------------
from reportlab.pdfgen import canvas as rl
def pdf_bytes(n=3):
    b = io.BytesIO(); c = rl.Canvas(b, pagesize=(300, 200))
    for i in range(1, n + 1):
        c.setFont('Helvetica-Bold', 20); c.drawString(30, 120, f'OmniView page {i}'); c.setFont('Helvetica', 12); c.drawString(30, 90, 'needle' if i == 2 else 'haystack'); c.showPage()
    c.save(); return b.getvalue()
mp = pdf_bytes(3); put('multipage.pdf', mp, 'pdf'); put('vector.ai', mp, 'pdf')  # PDF-compatible Illustrator file
from pypdf import PdfReader, PdfWriter
w = PdfWriter(); [w.add_page(p) for p in PdfReader(io.BytesIO(pdf_bytes(2))).pages]; w.encrypt('secret'); eb = io.BytesIO(); w.write(eb); put('encrypted.pdf', eb.getvalue(), 'pdf')
put('postscript.ai', '%!PS-Adobe-3.0\n%%Creator: Adobe Illustrator(R) 8.0\n%%BoundingBox: 0 0 100 100\n%%EndComments\nnewpath 0 0 moveto 100 100 lineto stroke\nshowpage\n', 'ai')
g = art(64, 48).convert('L'); rows = [''.join('%02X' % g.getpixel((x, y)) for x in range(64)) for y in range(48)]
PS_BODY = 'newpath 0 0 moveto 64 48 lineto stroke\nshowpage\n'
put('preview.eps', '%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 64 48\n%%BeginPreview: 64 48 8 48\n' + '\n'.join('% ' + r for r in rows) + '\n%%EndPreview\n%%EndComments\n' + PS_BODY, 'eps')
ps = ('%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 120 90\n%%EndComments\n' + PS_BODY).encode()
tiff = encode(art(120, 90), 'TIFF')
put('dos-tiff-preview.eps', b'\xc5\xd0\xd3\xc6' + struct.pack('<IIIIII', 30, len(ps), 0, 0, 30 + len(ps), len(tiff)) + b'\xff\xff' + ps + tiff, 'eps')
put('nopreview.eps', '%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 100 100\n%%EndComments\n' + PS_BODY, 'eps')

# ---- 3D (textured cube) ----------------------------------------------------------------------
tex = Image.new('RGB', (64, 64)); d = ImageDraw.Draw(tex)
for i in range(8):
    for j in range(8): d.rectangle([i * 8, j * 8, i * 8 + 7, j * 8 + 7], fill=(220, 60, 60) if (i + j) % 2 else (240, 240, 240))
tex_png = encode(tex, 'PNG')
FACES = [((1, 0, 0), [(1, -1, 1), (1, -1, -1), (1, 1, -1), (1, 1, 1)]), ((-1, 0, 0), [(-1, -1, -1), (-1, -1, 1), (-1, 1, 1), (-1, 1, -1)]),
         ((0, 1, 0), [(-1, 1, 1), (1, 1, 1), (1, 1, -1), (-1, 1, -1)]), ((0, -1, 0), [(-1, -1, -1), (1, -1, -1), (1, -1, 1), (-1, -1, 1)]),
         ((0, 0, 1), [(-1, -1, 1), (1, -1, 1), (1, 1, 1), (-1, 1, 1)]), ((0, 0, -1), [(1, -1, -1), (-1, -1, -1), (-1, 1, -1), (1, 1, -1)])]
UV = [(0, 0), (1, 0), (1, 1), (0, 1)]
pos = [c for _, cs in FACES for c in cs]; nor = [n for n, cs in FACES for _ in cs]; uvs = [UV[i] for _ in FACES for i in range(4)]
idx = [f * 4 + k for f in range(6) for k in (0, 1, 2, 0, 2, 3)]
def pad4(b): return b + b'\0' * (-len(b) % 4)
P = b''.join(struct.pack('<3f', *p) for p in pos); N = b''.join(struct.pack('<3f', *n) for n in nor); T = b''.join(struct.pack('<2f', *u) for u in uvs); I = pad4(b''.join(struct.pack('<H', i) for i in idx))
def gltf(embed):
    views = [P, N, T, I] + ([tex_png] if embed else [])
    blob = b''; bv = []
    for i, v in enumerate(views):
        bv.append({'buffer': 0, 'byteOffset': len(blob), 'byteLength': len(v), **({'target': 34963 if i == 3 else 34962} if i < 4 else {})}); blob += pad4(v)
    j = {'asset': {'version': '2.0', 'generator': 'omniview-fixtures'}, 'scene': 0, 'scenes': [{'nodes': [0]}], 'nodes': [{'mesh': 0}],
         'meshes': [{'primitives': [{'attributes': {'POSITION': 0, 'NORMAL': 1, 'TEXCOORD_0': 2}, 'indices': 3, 'material': 0}]}],
         'materials': [{'pbrMetallicRoughness': {'baseColorTexture': {'index': 0}, 'metallicFactor': 0}}], 'textures': [{'sampler': 0, 'source': 0}], 'samplers': [{}],
         'images': [{'bufferView': 4, 'mimeType': 'image/png'}] if embed else [{'uri': 'cube-texture.png'}],
         'accessors': [{'bufferView': 0, 'componentType': 5126, 'count': 24, 'type': 'VEC3', 'min': [-1, -1, -1], 'max': [1, 1, 1]}, {'bufferView': 1, 'componentType': 5126, 'count': 24, 'type': 'VEC3'},
                       {'bufferView': 2, 'componentType': 5126, 'count': 24, 'type': 'VEC2'}, {'bufferView': 3, 'componentType': 5123, 'count': 36, 'type': 'SCALAR'}],
         'bufferViews': bv, 'buffers': [{'byteLength': len(blob)} if embed else {'byteLength': len(blob), 'uri': 'cube.bin'}]}
    return j, blob
j, blob = gltf(True); js = json.dumps(j, separators=(',', ':')).encode(); js += b' ' * (-len(js) % 4)
put('cube.glb', b'glTF' + struct.pack('<II', 2, 12 + 8 + len(js) + 8 + len(blob)) + struct.pack('<I4s', len(js), b'JSON') + js + struct.pack('<I4s', len(blob), b'BIN\0') + blob, 'glb')
j, blob = gltf(False); put('cube.gltf', json.dumps(j, indent=1), 'gltf'); put('cube.bin', blob, 'bin'); put('cube-texture.png', tex_png, 'png')
obj = ['mtllib cube.mtl', 'o Cube'] + ['v %d %d %d' % p for p in pos] + ['vt %d %d' % u for u in UV] + ['vn %d %d %d' % n for n, _ in FACES] + ['usemtl Checker']
for f in range(6): obj.append('f ' + ' '.join(f'{f * 4 + k + 1}/{k + 1}/{f + 1}' for k in range(4)))
put('cube.obj', '\n'.join(obj) + '\n', 'obj'); put('cube.mtl', 'newmtl Checker\nKd 1 1 1\nmap_Kd cube-texture.png\n', 'mtl')

# ---- Office ----------------------------------------------------------------------------------
import docx, openpyxl
from pptx import Presentation
dd = docx.Document(); dd.add_heading('OmniView Word Fixture', 0); dd.add_paragraph('Hello from a real docx. The quick brown fox.')
tb = dd.add_table(rows=2, cols=2); tb.cell(0, 0).text = 'Cell A1'; tb.cell(1, 1).text = 'Cell B2'
with tempfile.TemporaryDirectory() as tmp:
    docx_path = os.path.join(tmp, 'sample.docx'); dd.save(docx_path)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Data'; ws.append(['id', 'name', 'value'])
    for i in range(1, 301): ws.append([i, f'row-{i}', i * 1.5])
    wb.create_sheet('Second')['A1'] = 'Hello Sheet'; xlsx_path = os.path.join(tmp, 'sample.xlsx'); wb.save(xlsx_path)
    pr = Presentation()
    for t in ('Slide One Title', 'Slide Two Title', 'Slide Three Title'):
        s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = t; s.placeholders[1].text = 'Body text for ' + t
    pptx_path = os.path.join(tmp, 'sample.pptx'); pr.save(pptx_path)
    for p, k in ((docx_path, 'docx'), (xlsx_path, 'xlsx'), (pptx_path, 'pptx')): put(os.path.basename(p), Path(p).read_bytes(), k)
    put('mislabeled-docx.txt', Path(docx_path).read_bytes(), 'docx')
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if soffice:
        env = dict(os.environ, HOME=tmp)
        for src, fmt, ext, kind in ((docx_path, 'odt', 'odt', 'odt'), (xlsx_path, 'ods', 'ods', 'ods'), (pptx_path, 'odp', 'odp', 'odp'), (docx_path, 'rtf', 'rtf', 'rtf'),
                                    (docx_path, 'doc', 'doc', 'legacy'), (pptx_path, 'ppt', 'ppt', 'legacy'), (xlsx_path, 'xls', 'xls', 'xlsx')):
            out = os.path.join(tmp, 'out_' + ext); os.makedirs(out, exist_ok=True)
            r = subprocess.run([soffice, '--headless', f'-env:UserInstallation=file://{tmp}/lo', '--convert-to', fmt, '--outdir', out, src], capture_output=True, env=env, timeout=180)
            f = Path(out) / ('sample.' + ext)
            if f.exists(): put('sample.' + ext, f.read_bytes(), kind)
            else: print('WARN: LibreOffice could not produce', ext, r.stderr[-200:], file=sys.stderr)
    else: print('WARN: soffice not found; ODF/RTF/legacy fixtures skipped', file=sys.stderr)
put('sample.csv', 'id,name,city\n1,Ada,London\n2,Grace,"New York"\n3,Linus,Helsinki\n', 'csv')
put('sample.txt', 'OmniView plain text fixture.\nSecond line.\n', 'text'); put('sample.md', '# Markdown fixture\n\n* item one\n* item two\n', 'text')

# ---- edge cases ------------------------------------------------------------------------------
put('corrupt.png', encode(art(100, 100), 'PNG')[:60], 'png')          # truncated: magic OK, data broken
put('mislabeled.jpg', encode(art(160, 120), 'PNG'), 'png')            # PNG content, .jpg name
(OUT / 'manifest.json').write_text(json.dumps(dict(sorted(EXPECT.items())), indent=1))
print(f'{len(EXPECT)} fixtures in {OUT}')
